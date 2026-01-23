"""
Gerador da planilha de escala 3 grupos
Replica EXATAMENTE o modelo MODELO_DE_ESCALA_DE_SOBREAVISO_3_GRUPOS.xlsx
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
import calendar
from pathlib import Path
from datetime import datetime


def gerar_planilha_escala_3_grupos(escala_id, mes, ano, prefeitura_nome, prefeitura_cidade):
    """
    Gera planilha de escala com 3 grupos
    Formato: MODELO_DE_ESCALA_DE_SOBREAVISO_3_GRUPOS.xlsx
    
    Args:
        escala_id: ID da escala
        mes: Mês (1-12)
        ano: Ano (YYYY)
        prefeitura_nome: Nome da prefeitura
        prefeitura_cidade: Cidade da prefeitura
        
    Returns:
        str: Caminho do arquivo gerado
    """
    from models.escala import Escala
    from models.escala_distribuicao import EscalaDistribuicao
    from models.servidor import Servidor
    
    # Buscar escala
    escala = Escala.query.get(escala_id)
    if not escala:
        raise ValueError(f'Escala {escala_id} não encontrada')
    
    # Buscar todas as distribuições
    distribuicoes = escala.distribuicoes.all()
    
    # Organizar servidores por grupo
    servidores_por_grupo = {
        'GERAL': {},
        'SAUDE_17H': {},
        'SAUDE_17H30': {}
    }
    
    for dist in distribuicoes:
        servidor = dist.servidor
        tipo_grupo = servidor.tipo_grupo
        
        if servidor.id not in servidores_por_grupo[tipo_grupo]:
            servidores_por_grupo[tipo_grupo][servidor.id] = {
                'servidor': servidor,
                'dias': []
            }
        
        servidores_por_grupo[tipo_grupo][servidor.id]['dias'].append(dist.dia_mes)
    
    # Criar workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Escala de Sobreaviso'
    
    # Configurar larguras de colunas
    ws.column_dimensions['A'].width = 10  # MAT.
    ws.column_dimensions['B'].width = 35  # NOME
    dias_no_mes = calendar.monthrange(ano, mes)[1]
    for col in range(3, dias_no_mes + 3):
        ws.column_dimensions[get_column_letter(col)].width = 4
    
    # Estilos
    fonte_titulo = Font(bold=True, size=12)
    fonte_header = Font(bold=True, size=10)
    fonte_normal = Font(size=10)
    
    alinhamento_centro = Alignment(horizontal='center', vertical='center', wrap_text=True)
    alinhamento_esquerda = Alignment(horizontal='left', vertical='center')
    
    borda_fina = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    preenchimento_header = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
    
    # Nome do mês em português
    meses = ['JANEIRO', 'FEVEREIRO', 'MARÇO', 'ABRIL', 'MAIO', 'JUNHO',
             'JULHO', 'AGOSTO', 'SETEMBRO', 'OUTUBRO', 'NOVEMBRO', 'DEZEMBRO']
    mes_nome = meses[mes - 1]
    
    linha_atual = 1
    
    # === CABEÇALHO PRINCIPAL ===
    ws.merge_cells(f'A{linha_atual}:{get_column_letter(dias_no_mes + 2)}{linha_atual}')
    ws[f'A{linha_atual}'] = 'RELATÓRIO DE ESCALA DE SOBREAVISO'
    ws[f'A{linha_atual}'].font = fonte_titulo
    ws[f'A{linha_atual}'].alignment = alinhamento_centro
    linha_atual += 1
    
    ws.merge_cells(f'A{linha_atual}:{get_column_letter(dias_no_mes + 2)}{linha_atual}')
    ws[f'A{linha_atual}'] = prefeitura_nome.upper()
    ws[f'A{linha_atual}'].font = fonte_normal
    ws[f'A{linha_atual}'].alignment = alinhamento_centro
    linha_atual += 1
    
    # Linha em branco
    linha_atual += 1
    
    ws.merge_cells(f'A{linha_atual}:{get_column_letter(dias_no_mes + 2)}{linha_atual}')
    ws[f'A{linha_atual}'] = f'LOCAL: {prefeitura_cidade}'
    ws[f'A{linha_atual}'].font = fonte_normal
    ws[f'A{linha_atual}'].alignment = alinhamento_esquerda
    linha_atual += 1
    
    ws.merge_cells(f'A{linha_atual}:{get_column_letter(dias_no_mes + 2)}{linha_atual}')
    ws[f'A{linha_atual}'] = f'COMPETÊNCIA: {mes_nome}/{ano}'
    ws[f'A{linha_atual}'].font = fonte_normal
    ws[f'A{linha_atual}'].alignment = alinhamento_esquerda
    linha_atual += 1
    
    # === GRUPO 1: GERAL ===
    ws.merge_cells(f'A{linha_atual}:{get_column_letter(dias_no_mes + 2)}{linha_atual}')
    ws[f'A{linha_atual}'] = 'HORÁRIO: 17h00 as 08h00 – EXCETO ÁREA DA SAÚDE'
    ws[f'A{linha_atual}'].font = fonte_normal
    ws[f'A{linha_atual}'].alignment = alinhamento_esquerda
    linha_atual += 1
    
    # Cabeçalho da tabela
    ws[f'A{linha_atual}'] = 'MAT.'
    ws[f'B{linha_atual}'] = 'NOME'
    for dia in range(1, dias_no_mes + 1):
        col_letra = get_column_letter(dia + 2)
        ws[f'{col_letra}{linha_atual}'] = dia
    
    for col in range(1, dias_no_mes + 3):
        col_letra = get_column_letter(col)
        cell = ws[f'{col_letra}{linha_atual}']
        cell.font = fonte_header
        cell.alignment = alinhamento_centro
        cell.border = borda_fina
        cell.fill = preenchimento_header
    
    linha_atual += 1
    linha_inicio_grupo1 = linha_atual
    
    # Dados dos servidores GERAL
    servidores_geral = sorted(servidores_por_grupo['GERAL'].values(), 
                              key=lambda x: x['servidor'].nome_completo)
    
    for srv_data in servidores_geral:
        servidor = srv_data['servidor']
        dias = srv_data['dias']
        
        ws[f'A{linha_atual}'] = servidor.matricula
        ws[f'A{linha_atual}'].alignment = alinhamento_centro
        ws[f'A{linha_atual}'].border = borda_fina
        
        ws[f'B{linha_atual}'] = servidor.nome_completo
        ws[f'B{linha_atual}'].alignment = alinhamento_esquerda
        ws[f'B{linha_atual}'].border = borda_fina
        
        for dia in range(1, dias_no_mes + 1):
            col_letra = get_column_letter(dia + 2)
            cell = ws[f'{col_letra}{linha_atual}']
            
            if dia in dias:
                cell.value = 'X'
                cell.alignment = alinhamento_centro
            
            cell.border = borda_fina
        
        linha_atual += 1
    
    # Linhas em branco
    linha_atual += 2
    
    # === GRUPO 2: SAÚDE 17h ===
    ws.merge_cells(f'A{linha_atual}:{get_column_letter(dias_no_mes + 2)}{linha_atual}')
    ws[f'A{linha_atual}'] = 'HORÁRIO: 17h00 as 08h00 (SAÚDE)'
    ws[f'A{linha_atual}'].font = fonte_normal
    ws[f'A{linha_atual}'].alignment = alinhamento_esquerda
    linha_atual += 1
    
    # Cabeçalho da tabela
    ws[f'A{linha_atual}'] = 'MAT.'
    ws[f'B{linha_atual}'] = 'NOME'
    for dia in range(1, dias_no_mes + 1):
        col_letra = get_column_letter(dia + 2)
        ws[f'{col_letra}{linha_atual}'] = dia
    
    for col in range(1, dias_no_mes + 3):
        col_letra = get_column_letter(col)
        cell = ws[f'{col_letra}{linha_atual}']
        cell.font = fonte_header
        cell.alignment = alinhamento_centro
        cell.border = borda_fina
        cell.fill = preenchimento_header
    
    linha_atual += 1
    
    # Dados dos servidores SAÚDE 17h
    servidores_saude17h = sorted(servidores_por_grupo['SAUDE_17H'].values(),
                                 key=lambda x: x['servidor'].nome_completo)
    
    for srv_data in servidores_saude17h:
        servidor = srv_data['servidor']
        dias = srv_data['dias']
        
        ws[f'A{linha_atual}'] = servidor.matricula
        ws[f'A{linha_atual}'].alignment = alinhamento_centro
        ws[f'A{linha_atual}'].border = borda_fina
        
        ws[f'B{linha_atual}'] = servidor.nome_completo
        ws[f'B{linha_atual}'].alignment = alinhamento_esquerda
        ws[f'B{linha_atual}'].border = borda_fina
        
        for dia in range(1, dias_no_mes + 1):
            col_letra = get_column_letter(dia + 2)
            cell = ws[f'{col_letra}{linha_atual}']
            
            if dia in dias:
                cell.value = 'X'
                cell.alignment = alinhamento_centro
            
            cell.border = borda_fina
        
        linha_atual += 1
    
    # Linhas em branco
    linha_atual += 2
    
    # === GRUPO 3: SAÚDE 17h30 ===
    ws.merge_cells(f'A{linha_atual}:{get_column_letter(dias_no_mes + 2)}{linha_atual}')
    ws[f'A{linha_atual}'] = 'HORÁRIO: 17h30 as 08h00 (SAÚDE)'
    ws[f'A{linha_atual}'].font = fonte_normal
    ws[f'A{linha_atual}'].alignment = alinhamento_esquerda
    linha_atual += 1
    
    # Cabeçalho da tabela
    ws[f'A{linha_atual}'] = 'MAT.'
    ws[f'B{linha_atual}'] = 'NOME'
    for dia in range(1, dias_no_mes + 1):
        col_letra = get_column_letter(dia + 2)
        ws[f'{col_letra}{linha_atual}'] = dia
    
    for col in range(1, dias_no_mes + 3):
        col_letra = get_column_letter(col)
        cell = ws[f'{col_letra}{linha_atual}']
        cell.font = fonte_header
        cell.alignment = alinhamento_centro
        cell.border = borda_fina
        cell.fill = preenchimento_header
    
    linha_atual += 1
    
    # Dados dos servidores SAÚDE 17h30
    servidores_saude1730 = sorted(servidores_por_grupo['SAUDE_17H30'].values(),
                                  key=lambda x: x['servidor'].nome_completo)
    
    for srv_data in servidores_saude1730:
        servidor = srv_data['servidor']
        dias = srv_data['dias']
        
        ws[f'A{linha_atual}'] = servidor.matricula
        ws[f'A{linha_atual}'].alignment = alinhamento_centro
        ws[f'A{linha_atual}'].border = borda_fina
        
        ws[f'B{linha_atual}'] = servidor.nome_completo
        ws[f'B{linha_atual}'].alignment = alinhamento_esquerda
        ws[f'B{linha_atual}'].border = borda_fina
        
        for dia in range(1, dias_no_mes + 1):
            col_letra = get_column_letter(dia + 2)
            cell = ws[f'{col_letra}{linha_atual}']
            
            if dia in dias:
                cell.value = 'X'
                cell.alignment = alinhamento_centro
            
            cell.border = borda_fina
        
        linha_atual += 1
    
    # Salvar arquivo
    from flask import current_app
    export_folder = Path(current_app.config['EXPORT_FOLDER'])
    export_folder.mkdir(parents=True, exist_ok=True)
    
    nome_arquivo = f'Escala_3_Grupos_{mes:02d}_{ano}.xlsx'
    caminho_arquivo = export_folder / nome_arquivo
    
    wb.save(caminho_arquivo)
    
    return str(caminho_arquivo)