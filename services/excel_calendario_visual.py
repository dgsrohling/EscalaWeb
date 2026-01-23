"""
Gerador do calendário visual
Replica EXATAMENTE o modelo MODELO_DE_CALENDARIO_DE_ESCALA_GRUPO_1.xlsx
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import calendar
from pathlib import Path
from datetime import date


def gerar_planilha_calendario_visual(escala_id, mes, ano):
    """
    Gera planilha de calendário visual mensal (apenas grupo GERAL)
    Formato: MODELO_DE_CALENDARIO_DE_ESCALA_GRUPO_1.xlsx
    
    Args:
        escala_id: ID da escala
        mes: Mês (1-12)
        ano: Ano (YYYY)
        
    Returns:
        str: Caminho do arquivo gerado
    """
    from models.escala import Escala
    
    # Buscar escala
    escala = Escala.query.get(escala_id)
    if not escala:
        raise ValueError(f'Escala {escala_id} não encontrada')
    
    # Buscar distribuições apenas do grupo GERAL
    distribuicoes_dict = {}
    for dist in escala.distribuicoes:
        if dist.servidor.tipo_grupo == 'GERAL':
            distribuicoes_dict[dist.dia_mes] = dist.servidor.nome_completo.split()[0]  # Primeiro nome
    
    # Criar workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Calendário'
    
    # Configurar larguras de colunas (7 colunas para 7 dias da semana)
    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = 18
    
    # Estilos
    fonte_titulo = Font(bold=True, size=14)
    fonte_dia_semana = Font(bold=True, size=11)
    fonte_dia = Font(size=10)
    fonte_nome = Font(size=9)
    
    alinhamento_centro = Alignment(horizontal='center', vertical='center')
    alinhamento_topo = Alignment(horizontal='center', vertical='top')
    
    borda_fina = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Nome do mês em português
    meses = ['JANEIRO', 'FEVEREIRO', 'MARÇO', 'ABRIL', 'MAIO', 'JUNHO',
             'JULHO', 'AGOSTO', 'SETEMBRO', 'OUTUBRO', 'NOVEMBRO', 'DEZEMBRO']
    mes_nome = meses[mes - 1]
    
    # === CABEÇALHO ===
    linha_atual = 1
    
    ws.merge_cells(f'A{linha_atual}:G{linha_atual}')
    ws[f'A{linha_atual}'] = f'{mes_nome} – {ano}'
    ws[f'A{linha_atual}'].font = fonte_titulo
    ws[f'A{linha_atual}'].alignment = alinhamento_centro
    linha_atual += 1
    
    # === DIAS DA SEMANA ===
    dias_semana = ['Domingo', 'Segunda', 'Terça', 'Quarta', 'Quinta', 'Sexta', 'Sábado']
    for col, dia_semana in enumerate(dias_semana, start=1):
        cell = ws.cell(row=linha_atual, column=col)
        cell.value = dia_semana
        cell.font = fonte_dia_semana
        cell.alignment = alinhamento_centro
        cell.border = borda_fina
    
    linha_atual += 1
    
    # === CALENDÁRIO ===
    # Obter informações do mês
    primeiro_dia_semana = date(ano, mes, 1).weekday()
    # weekday() retorna 0=segunda, 6=domingo
    # Converter para 0=domingo, 6=sábado
    primeiro_dia_semana = (primeiro_dia_semana + 1) % 7
    
    dias_no_mes = calendar.monthrange(ano, mes)[1]
    
    # Calcular quantas semanas precisamos (linhas)
    total_celulas = primeiro_dia_semana + dias_no_mes
    num_semanas = (total_celulas + 6) // 7  # Arredondar para cima
    
    # Preencher calendário
    dia_atual = 1
    
    for semana in range(num_semanas):
        # Cada semana ocupa 2 linhas: uma para o dia e outra para o nome
        linha_dia = linha_atual
        linha_nome = linha_atual + 1
        
        for dia_semana in range(7):
            col = dia_semana + 1
            
            # Primeira semana - pular dias antes do início do mês
            if semana == 0 and dia_semana < primeiro_dia_semana:
                cell_dia = ws.cell(row=linha_dia, column=col)
                cell_dia.border = borda_fina
                
                cell_nome = ws.cell(row=linha_nome, column=col)
                cell_nome.border = borda_fina
                continue
            
            # Se ainda há dias no mês
            if dia_atual <= dias_no_mes:
                # Célula do dia
                cell_dia = ws.cell(row=linha_dia, column=col)
                cell_dia.value = dia_atual
                cell_dia.font = fonte_dia
                cell_dia.alignment = alinhamento_centro
                cell_dia.border = borda_fina
                
                # Célula do nome do plantonista
                cell_nome = ws.cell(row=linha_nome, column=col)
                if dia_atual in distribuicoes_dict:
                    cell_nome.value = distribuicoes_dict[dia_atual]
                    cell_nome.font = fonte_nome
                    cell_nome.alignment = alinhamento_centro
                cell_nome.border = borda_fina
                
                dia_atual += 1
            else:
                # Células vazias após o último dia do mês
                cell_dia = ws.cell(row=linha_dia, column=col)
                cell_dia.border = borda_fina
                
                cell_nome = ws.cell(row=linha_nome, column=col)
                cell_nome.border = borda_fina
        
        # Definir altura das linhas
        ws.row_dimensions[linha_dia].height = 20
        ws.row_dimensions[linha_nome].height = 25
        
        linha_atual += 2  # Avançar 2 linhas (dia + nome)
    
    # Salvar arquivo
    from flask import current_app
    export_folder = Path(current_app.config['EXPORT_FOLDER'])
    export_folder.mkdir(parents=True, exist_ok=True)
    
    nome_arquivo = f'Calendario_Visual_{mes:02d}_{ano}.xlsx'
    caminho_arquivo = export_folder / nome_arquivo
    
    wb.save(caminho_arquivo)
    
    return str(caminho_arquivo)