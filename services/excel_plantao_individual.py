"""
Gerador de planilha individual de plantões
Replica EXATAMENTE o modelo JULIANO_34871_12-2025.xls
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
import calendar
from pathlib import Path


def gerar_planilha_plantao_individual(servidor_id, mes, ano):
    """
    Gera planilha individual de plantões realizados
    Formato: NOME_MATRICULA_MES-ANO.xlsx
    
    Args:
        servidor_id: ID do servidor
        mes: Mês (1-12)
        ano: Ano (YYYY)
        
    Returns:
        str: Caminho do arquivo gerado
    """
    from models.servidor import Servidor
    from models.plantao_realizado import PlantaoRealizado
    
    # Buscar servidor
    servidor = Servidor.query.get(servidor_id)
    if not servidor:
        raise ValueError(f'Servidor {servidor_id} não encontrado')
    
    # Buscar plantões realizados
    plantoes = PlantaoRealizado.buscar_por_servidor_e_competencia(servidor_id, mes, ano)
    
    # Criar dicionário de plantões por dia
    plantoes_dict = {p.dia_mes: p for p in plantoes}
    
    # Criar workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Planilha1'
    
    # Configurar larguras de colunas
    ws.column_dimensions['A'].width = 8   # DIAS
    ws.column_dimensions['B'].width = 18  # HORÁRIO INÍCIO
    ws.column_dimensions['C'].width = 18  # HORÁRIO FINAL
    ws.column_dimensions['D'].width = 15  # HORAS EXTRAS
    ws.column_dimensions['E'].width = 15  # ADC. NOTURNO
    
    # Estilos
    fonte_titulo = Font(bold=True, size=12)
    fonte_header = Font(bold=True, size=10)
    fonte_normal = Font(size=10)
    
    alinhamento_centro = Alignment(horizontal='center', vertical='center')
    alinhamento_esquerda = Alignment(horizontal='left', vertical='center')
    
    borda_fina = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    preenchimento_header = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
    
    # Nome do mês em português
    meses = ['JANEIRO', 'FEVEREIRO', 'MARÇO', 'ABRIL', 'MAIO', 'JUNHO',
             'JULHO', 'AGOSTO', 'SETEMBRO', 'OUTUBRO', 'NOVEMBRO', 'DEZEMBRO']
    mes_nome = meses[mes - 1].upper()
    
    # === CABEÇALHO ===
    linha_atual = 1
    
    # Linha 1: Título
    ws.merge_cells(f'A{linha_atual}:E{linha_atual}')
    ws[f'A{linha_atual}'] = f'RELATÓRIO DE HORAS EXTRAS DE SOBREAVISO – {mes_nome}/{ano}'
    ws[f'A{linha_atual}'].font = fonte_titulo
    ws[f'A{linha_atual}'].alignment = alinhamento_centro
    ws[f'A{linha_atual}'].fill = preenchimento_header
    linha_atual += 1
    
    # Linha 2: Secretaria/Local
    ws.merge_cells(f'A{linha_atual}:E{linha_atual}')
    if servidor.secretaria:
        local = servidor.secretaria.nome
    else:
        local = 'GERAL'
    ws[f'A{linha_atual}'] = local
    ws[f'A{linha_atual}'].font = fonte_normal
    ws[f'A{linha_atual}'].alignment = alinhamento_esquerda
    linha_atual += 1
    
    # Linha 3: Servidor e Matrícula
    ws.merge_cells(f'A{linha_atual}:C{linha_atual}')
    ws[f'A{linha_atual}'] = f'SERVIDOR: {servidor.nome_completo}'
    ws[f'A{linha_atual}'].font = fonte_normal
    ws[f'A{linha_atual}'].alignment = alinhamento_esquerda
    
    ws[f'D{linha_atual}'] = 'MATRÍCULA:'
    ws[f'D{linha_atual}'].font = fonte_normal
    ws[f'D{linha_atual}'].alignment = alinhamento_esquerda
    
    ws[f'E{linha_atual}'] = servidor.matricula
    ws[f'E{linha_atual}'].font = fonte_normal
    ws[f'E{linha_atual}'].alignment = alinhamento_centro
    linha_atual += 1
    
    # === CABEÇALHO DA TABELA ===
    headers = ['DIAS', 'HORÁRIO INÍCIO', 'HORÁRIO FINAL', 'HORAS EXTRAS', 'ADC. NOTURNO']
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=linha_atual, column=col_idx)
        cell.value = header
        cell.font = fonte_header
        cell.alignment = alinhamento_centro
        cell.border = borda_fina
        cell.fill = preenchimento_header
    
    linha_atual += 1
    
    # === DADOS DOS DIAS ===
    dias_no_mes = calendar.monthrange(ano, mes)[1]
    
    total_horas_extras = 0
    total_adicional_noturno = 0
    
    for dia in range(1, dias_no_mes + 1):
        # Coluna A: Dia
        ws.cell(row=linha_atual, column=1).value = dia
        ws.cell(row=linha_atual, column=1).alignment = alinhamento_centro
        ws.cell(row=linha_atual, column=1).border = borda_fina
        
        # Verificar se há plantão neste dia
        if dia in plantoes_dict:
            plantao = plantoes_dict[dia]
            
            # Coluna B: Horário Início
            ws.cell(row=linha_atual, column=2).value = plantao.hora_inicio.strftime('%H:%M') if plantao.hora_inicio else ''
            ws.cell(row=linha_atual, column=2).alignment = alinhamento_centro
            ws.cell(row=linha_atual, column=2).border = borda_fina
            
            # Coluna C: Horário Final
            ws.cell(row=linha_atual, column=3).value = plantao.hora_fim.strftime('%H:%M') if plantao.hora_fim else ''
            ws.cell(row=linha_atual, column=3).alignment = alinhamento_centro
            ws.cell(row=linha_atual, column=3).border = borda_fina
            
            # Coluna D: Horas Extras
            horas_extras = plantao.horas_trabalhadas or 0
            ws.cell(row=linha_atual, column=4).value = PlantaoRealizado.get_horas_display(horas_extras)
            ws.cell(row=linha_atual, column=4).alignment = alinhamento_centro
            ws.cell(row=linha_atual, column=4).border = borda_fina
            
            # Coluna E: Adicional Noturno
            adicional = plantao.adicional_noturno or 0
            ws.cell(row=linha_atual, column=5).value = PlantaoRealizado.get_horas_display(adicional)
            ws.cell(row=linha_atual, column=5).alignment = alinhamento_centro
            ws.cell(row=linha_atual, column=5).border = borda_fina
            
            total_horas_extras += horas_extras
            total_adicional_noturno += adicional
        else:
            # Dia sem plantão - células vazias com borda
            for col in range(2, 6):
                ws.cell(row=linha_atual, column=col).border = borda_fina
        
        linha_atual += 1
    
    # === LINHA DE TOTAL ===
    ws.cell(row=linha_atual, column=1).value = 'TOTAL'
    ws.cell(row=linha_atual, column=1).font = fonte_header
    ws.cell(row=linha_atual, column=1).alignment = alinhamento_centro
    ws.cell(row=linha_atual, column=1).border = borda_fina
    ws.cell(row=linha_atual, column=1).fill = preenchimento_header
    
    # Células vazias
    for col in [2, 3]:
        ws.cell(row=linha_atual, column=col).border = borda_fina
        ws.cell(row=linha_atual, column=col).fill = preenchimento_header
    
    # Total Horas Extras
    ws.cell(row=linha_atual, column=4).value = PlantaoRealizado.get_horas_display(total_horas_extras)
    ws.cell(row=linha_atual, column=4).font = fonte_header
    ws.cell(row=linha_atual, column=4).alignment = alinhamento_centro
    ws.cell(row=linha_atual, column=4).border = borda_fina
    ws.cell(row=linha_atual, column=4).fill = preenchimento_header
    
    # Total Adicional Noturno
    ws.cell(row=linha_atual, column=5).value = PlantaoRealizado.get_horas_display(total_adicional_noturno)
    ws.cell(row=linha_atual, column=5).font = fonte_header
    ws.cell(row=linha_atual, column=5).alignment = alinhamento_centro
    ws.cell(row=linha_atual, column=5).border = borda_fina
    ws.cell(row=linha_atual, column=5).fill = preenchimento_header
    
    # Salvar arquivo
    from flask import current_app
    export_folder = Path(current_app.config['EXPORT_FOLDER'])
    export_folder.mkdir(parents=True, exist_ok=True)
    
    # Nome do arquivo: NOME_MATRICULA_MES-ANO.xlsx
    nome_limpo = servidor.nome_completo.replace(' ', '_')
    nome_arquivo = f'{nome_limpo}_{servidor.matricula}_{mes:02d}-{ano}.xlsx'
    caminho_arquivo = export_folder / nome_arquivo
    
    wb.save(caminho_arquivo)
    
    return str(caminho_arquivo)